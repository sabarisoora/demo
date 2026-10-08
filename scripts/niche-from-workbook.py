#!/usr/bin/env python3
"""
Create a ProfitIQS niche from an Essential workbook (the PROFITIQS Excel product).

  pip install openpyxl
  python3 scripts/niche-from-workbook.py "path/to/<Niche>_Essential_2026.xlsx" [--slug hair-salon]

Writes:
  src/niches/<slug>.ts                 niche config (labels, streams, categories, thresholds, copy)
  src/niches/samples/<slug>.json       the workbook's sample rows
  src/niches/expected/<slug>.json      Excel's cached totals; the test suite checks the app matches them
and registers the niche in src/niches/index.ts. Its landing page is then live at /<slug>.

Review the generated .ts file: landing copy, segment thresholds and default goals are starting points.
Industry benchmarks are left out on purpose; add them only from a real source.
"""
import argparse
import json
import re
import sys
from datetime import date, datetime
from pathlib import Path

try:
    import openpyxl
except ImportError:
    sys.exit("openpyxl is required: pip install openpyxl")

ROOT = Path(__file__).resolve().parent.parent
NICHES = ROOT / "src" / "niches"


def die(msg):
    sys.exit(f"✗ {msg}")


def slugify(s):
    return re.sub(r"[^a-z0-9]+", "-", s.lower()).strip("-")


def title_case(s):
    small = {"and", "or", "of", "the", "a", "an", "to", "for", "in", "on"}
    words = s.lower().split()
    return " ".join(w if (i and w in small) else w.capitalize() for i, w in enumerate(words))


def iso(v):
    if isinstance(v, datetime):
        return v.date().isoformat()
    if isinstance(v, date):
        return v.isoformat()
    return str(v)[:10]


def dv_list(ws, col_letter):
    """The dropdown list (data validation) applied to a column, if any."""
    for dv in ws.data_validations.dataValidation:
        if any(str(r).startswith(col_letter) for r in str(dv.sqref).split()) and dv.formula1:
            f = dv.formula1.strip('"')
            if "," in f:
                return [x.strip() for x in f.split(",") if x.strip()]
    return []


def header_row(ws):
    for r in range(1, 11):
        vals = [str(c.value or "").strip() for c in ws[r]]
        if any(v.lower() == "date" for v in vals):
            return r, vals
    die(f"no header row with a 'Date' column in sheet '{ws.title}'")


def num(v):
    return round(float(v), 2) if isinstance(v, (int, float)) else 0.0


def ts_str(s):
    return json.dumps(s, ensure_ascii=False)


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("workbook")
    ap.add_argument("--slug", help="URL slug, e.g. hair-salon (default: from the workbook name)")
    ap.add_argument("--out", default=str(NICHES), help="output folder (default: src/niches)")
    ap.add_argument("--no-register", action="store_true", help="don't edit src/niches/index.ts")
    a = ap.parse_args()

    wb = openpyxl.load_workbook(a.workbook)
    cached = openpyxl.load_workbook(a.workbook, data_only=True)
    sheets = {ws.title.upper(): ws for ws in wb.worksheets}

    # ── Niche name: "Auto Repair Shop Business Intelligence System™"
    start = sheets.get("START HERE")
    title = str(start["B2"].value or "") if start else ""
    m = re.match(r"(.+?)\s+Business Intelligence System", title)
    if not m:
        die(f"couldn't read the niche name from START HERE!B2 ({title!r})")
    name = m.group(1).strip()
    slug = a.slug or slugify(name)
    var = re.sub(r"-(\w)", lambda x: x.group(1).upper(), slug)
    noun = name.split()[-1].lower()

    # ── The job entry sheet: "<SOMETHING> ENTRY", not the expense one
    entry = next((ws for t, ws in sheets.items() if t.endswith("ENTRY") and "EXPENSE" not in t), None)
    expense = sheets.get("EXPENSE ENTRY")
    if not entry or not expense:
        die("need a '<job> ENTRY' sheet and an 'EXPENSE ENTRY' sheet")
    singular = title_case(re.sub(r"\s+ENTRY$", "", entry.title.upper()))
    plural = singular + ("es" if singular.endswith(("s", "x", "ch", "sh")) else "s")

    hr, headers = header_row(entry)
    lower = [h.lower() for h in headers]
    i_date = lower.index("date")
    pairs = []
    for i, h in enumerate(headers):
        mm = re.match(r"(.+?)\s+revenue$", h, re.I)
        if mm:
            base = mm.group(1)
            j = next((k for k, x in enumerate(headers) if x.lower() == f"{base.lower()} cost"), None)
            if j is None:
                die(f"found '{h}' but no matching '{base} Cost' column")
            pairs.append((base.strip(), i, j))
    if len(pairs) != 2:
        die(f"expected exactly two '<X> Revenue' / '<X> Cost' column pairs, found {len(pairs)}: {[p[0] for p in pairs]}")
    (a_name, a_rev, a_cost), (b_name, b_rev, b_cost) = pairs
    # Text columns between the date and the first amount: category, then customer.
    text_cols = [i for i in range(i_date + 1, min(a_rev, b_rev))]
    if len(text_cols) < 2:
        die("expected a category column and a customer column after Date")
    i_cat, i_cust = text_cols[0], text_cols[1]
    customer_label = headers[i_cust] or "Customer"

    col = lambda i: openpyxl.utils.get_column_letter(i + 1)
    categories = dv_list(entry, col(i_cat))
    exp_hr, exp_headers = header_row(expense)
    exp_lower = [h.lower() for h in exp_headers]
    e_date, e_cat = exp_lower.index("date"), exp_lower.index("category")
    e_vendor = next((i for i, h in enumerate(exp_lower) if h in ("vendor", "payee", "supplier")), e_cat + 1)
    e_amount = exp_lower.index("amount")
    exp_categories = dv_list(expense, col(e_cat))

    # ── Sample rows
    orders, prefix = [], None
    for r in entry.iter_rows(min_row=hr + 1, values_only=True):
        if not r[0] or not r[i_date]:
            continue
        prefix = prefix or re.match(r"[A-Za-z]+-?", str(r[0])).group(0)
        orders.append({
            "ref": str(r[0]), "date": iso(r[i_date]), "category": str(r[i_cat] or ""), "customer": str(r[i_cust] or ""),
            "revenueA": num(r[a_rev]), "costA": num(r[a_cost]), "revenueB": num(r[b_rev]), "costB": num(r[b_cost]),
        })
    exps = []
    for r in expense.iter_rows(min_row=exp_hr + 1, values_only=True):
        if not r[0] or not r[e_date]:
            continue
        exps.append({"ref": str(r[0]), "date": iso(r[e_date]), "category": str(r[e_cat] or ""), "vendor": str(r[e_vendor] or ""), "amount": num(r[e_amount])})
    if not orders:
        die("no sample rows found")
    categories = categories or sorted({o["category"] for o in orders})
    exp_categories = exp_categories or sorted({e["category"] for e in exps})
    if "Other" not in exp_categories:
        exp_categories.append("Other")

    # ── Planning threshold from the dashboard insight ("below the 15% ... threshold — <tip>")
    threshold, tip = 0.15, f"Review pricing and costs on your biggest {plural.lower()}."
    dash = sheets.get("DASHBOARD")
    if dash:
        for row in dash.iter_rows():
            for c in row:
                v = str(c.value or "")
                mm = re.search(r"below the (\d+(?:\.\d+)?)% [^\"]*?threshold\s*[—-]\s*([^\"]+?)\"", v)
                if mm:
                    threshold = float(mm.group(1)) / 100
                    # Drop spreadsheet references like "in REPAIR ORDER ENTRY"; the app has no sheets.
                    t = re.sub(r"\s+in\s+[A-Z][A-Z &/]+$", "", mm.group(2).strip().rstrip("."))
                    tip = t[0].upper() + t[1:] + "."

    # ── Excel's cached totals for the test suite (TAX & DEDUCTIONS sheet)
    tax = cached.worksheets[[w.title.upper() for w in cached.worksheets].index("TAX & DEDUCTIONS")] if "TAX & DEDUCTIONS" in sheets else None
    expected = None
    if tax is not None and isinstance(tax["B5"].value, (int, float)):
        expected = {k: tax[c].value for k, c in [("revenue", "B5"), ("directCost", "B6"), ("overhead", "B7"), ("netProfit", "B9")]}

    # ── Derived defaults
    revenue = sum(o["revenueA"] + o["revenueB"] + 0 for o in orders)
    avg = revenue / len(orders)
    round_to = lambda x, step: max(step, round(x / step) * step)
    goals = {
        "revenue": round_to(revenue * 1.25, 10000),
        "netProfit": round_to(revenue * 1.25 * max(threshold, 0.1), 1000),
        "avgTicket": round_to(avg * 1.1, 10),
        "jobCount": round_to(len(orders) * 1.25, 10),
        "grossMargin": 0.55,
        "netMargin": threshold,
    }
    segments = {"vip": round_to(avg * 2.5, 50), "core": round_to(avg, 50)}

    out = Path(a.out)
    (out / "samples").mkdir(parents=True, exist_ok=True)
    (out / "expected").mkdir(parents=True, exist_ok=True)
    (out / "samples" / f"{slug}.json").write_text(json.dumps({"orders": orders, "expenses": exps}))
    if expected:
        (out / "expected" / f"{slug}.json").write_text(json.dumps(expected, indent=2) + "\n")

    short = (prefix or singular[:2]).rstrip("-").upper()
    lst = lambda xs: "[\n" + "".join(f"    {ts_str(x)},\n" for x in xs) + "  ]"
    ts = f'''import type {{ Niche, SampleData }} from "./types";
import sample from "./samples/{slug}.json";

// Generated from "{Path(a.workbook).name}" by scripts/niche-from-workbook.py.
// Review the copy, segment thresholds and default goals before launch.
export const {var}: Niche = {{
  slug: {ts_str(slug)},
  name: {ts_str(name)},
  businessNoun: {ts_str(noun)},
  job: {{
    singular: {ts_str(singular)},
    plural: {ts_str(plural)},
    short: {ts_str(short)},
    refPrefix: {ts_str(prefix or short + "-")},
    customerLabel: {ts_str(customer_label)},
  }},
  streams: {{ a: {ts_str(a_name)}, b: {ts_str(b_name)} }},
  details: {{
    technician: "Team member",
    technicianPlural: "Team",
    hours: "Hours",
    comeback: "Redo",
    comebackHint: "Work that had to be redone",
  }},
  jobCategories: {lst(categories)},
  expenseCategories: {lst(exp_categories)},
  thresholds: {{
    netMargin: {threshold},
    jobGrossMargin: 0.4,
    streamAMargin: 0.4,
    streamBMargin: 0.5,
    overheadShare: 0.3,
  }},
  // Lifetime revenue: VIP ≈ 2.5× and Core ≈ 1× the sample's average ticket.
  segments: {{ vip: {segments["vip"]}, core: {segments["core"]} }},
  defaultGoals: {json.dumps(goals)},
  copy: {{
    lowMarginTip: {ts_str(tip)},
    heroTitle: {ts_str(f"Know exactly what your {noun} made — and what to set aside for tax.")},
    heroSubtitle: {ts_str(f"Log {plural.lower()} and expenses (or import a CSV). ProfitIQS turns them into true profit per {singular.lower()}, a tax reserve, and a plain-English health check.")},
    pains: [
      {{ title: "Busy, but where's the profit?", body: {ts_str(f"See net profit after {a_name.lower()}, {b_name.lower()} and overhead, not just sales.")} }},
      {{ title: "Tax bill surprises", body: "A recommended tax reserve for your country, how much you've funded, and your VAT/GST position." }},
      {{ title: {ts_str(f"Which {plural.lower()} actually pay")}, body: {ts_str(f"Profit per {singular.lower()} and per category, so you know what to push and what to reprice.")} }},
    ],
  }},
  sample: sample as SampleData,
}};
'''
    (out / f"{slug}.ts").write_text(ts)

    if not a.no_register:
        idx = NICHES / "index.ts"
        s = idx.read_text()
        if f'from "./{slug}"' not in s:
            s = s.replace('import type { Niche } from "./types";', f'import {{ {var} }} from "./{slug}";\nimport type {{ Niche }} from "./types";')
            s = s.replace("  [autoRepair.slug]: autoRepair,", f"  [autoRepair.slug]: autoRepair,\n  [{var}.slug]: {var},")
            idx.write_text(s)

    print(f"✓ {name}: {len(orders)} {plural.lower()}, {len(exps)} expenses, streams {a_name}/{b_name}, {len(categories)} categories")
    print(f"  wrote {out / (slug + '.ts')}" + ("" if a.no_register else " and registered it"))
    print(f"  landing page: /{slug}   signup: /signup?niche={slug}")
    if not expected:
        print("  ! no cached Excel totals found (open and save the file in Excel, then re-run to enable the match test)")


if __name__ == "__main__":
    main()
